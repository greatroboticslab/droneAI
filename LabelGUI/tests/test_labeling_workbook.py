"""The queue and its downloadable Excel copy must tell the same story."""
import io
import sys
import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook, load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from db.db_store import DBStore  # noqa: E402


def source_workbook():
    book = Workbook()
    sheet = book.active
    sheet.append(["Persona Name", "Youtube Link", "Notes"])
    sheet.append(["Ann", "https://youtu.be/aaaaaaaaaaa", "keep"])
    sheet.append(["Bob", "https://youtu.be/bbbbbbbbbbb", "remove"])
    stream = io.BytesIO()
    book.save(stream)
    return stream.getvalue()


class LabelingWorkbookTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.db = DBStore(str(Path(self.tmp.name) / "test.sqlite"))
        self.db.save_dataset("ds", "Flights", "flights.xlsx", source_workbook(), "tester")
        self.db.replace_dataset_items("ds", [
            {"item_key": "ds:1", "row_index": 1, "person_name": "Ann",
             "youtube_link": "https://youtu.be/aaaaaaaaaaa"},
            {"item_key": "ds:2", "row_index": 2, "person_name": "Bob",
             "youtube_link": "https://youtu.be/bbbbbbbbbbb"},
        ])

    def test_completion_updates_queue_and_excel_status(self):
        self.db.complete_dataset_item("ds:1", "labeler", "Real flight")
        item = self.db.get_dataset_item("ds:1")
        self.assertEqual(item["status"], "labeled")
        self.assertEqual(item["labeled_by"], "labeler")
        sheet = load_workbook(io.BytesIO(self.db.get_dataset_file_blob("ds"))).active
        headers = [cell.value for cell in sheet[1]]
        self.assertEqual(sheet.cell(2, headers.index("Label Status") + 1).value, "Labeled")
        self.assertEqual(sheet.cell(2, headers.index("Labeled By") + 1).value, "labeler")
        self.assertEqual(sheet["C2"].value, "keep")

    def test_unavailable_removal_deletes_only_matching_excel_row_and_queue_item(self):
        self.db.remove_unavailable_dataset_item("ds:2")
        self.assertIsNone(self.db.get_dataset_item("ds:2"))
        self.assertEqual(self.db.dataset_stats("ds")["total"], 1)
        sheet = load_workbook(io.BytesIO(self.db.get_dataset_file_blob("ds"))).active
        self.assertEqual(sheet.max_row, 2)
        self.assertEqual(sheet["A2"].value, "Ann")
        self.assertEqual(sheet["C2"].value, "keep")

    def test_missing_excel_match_does_not_delete_queue_item(self):
        book = Workbook()
        book.active.append(["Persona Name", "Youtube Link"])
        book.active.append(["Ann", "https://youtu.be/aaaaaaaaaaa"])
        stream = io.BytesIO()
        book.save(stream)
        self.db.save_dataset("ds", "Flights", "flights.xlsx", stream.getvalue(), "tester")
        with self.assertRaises(ValueError):
            self.db.remove_unavailable_dataset_item("ds:2")
        self.assertIsNotNone(self.db.get_dataset_item("ds:2"))

    def test_reset_removes_labeled_state_from_excel_and_queue(self):
        self.db.complete_dataset_item("ds:1", "labeler", "Real flight")
        self.db.reset_dataset_item("ds:1")
        item = self.db.get_dataset_item("ds:1")
        self.assertEqual(item["status"], "not_labeled")
        self.assertEqual(item["labeled_by"], "")
        sheet = load_workbook(io.BytesIO(self.db.get_dataset_file_blob("ds"))).active
        headers = [cell.value for cell in sheet[1]]
        self.assertIsNone(sheet.cell(2, headers.index("Label Status") + 1).value)
        self.assertIsNone(sheet.cell(2, headers.index("Labeled By") + 1).value)

    def test_duplicate_links_remove_the_specific_queue_row(self):
        book = Workbook()
        book.active.append(["Persona Name", "Youtube Link", "Notes"])
        book.active.append(["Ann", "https://youtu.be/aaaaaaaaaaa", "first"])
        book.active.append(["Ann", "https://youtu.be/aaaaaaaaaaa", "second"])
        stream = io.BytesIO()
        book.save(stream)
        self.db.save_dataset("ds", "Flights", "flights.xlsx", stream.getvalue(), "tester")
        self.db.replace_dataset_items("ds", [
            {"item_key": "ds:1", "row_index": 1, "person_name": "Ann",
             "youtube_link": "https://youtu.be/aaaaaaaaaaa"},
            {"item_key": "ds:2", "row_index": 2, "person_name": "Ann",
             "youtube_link": "https://youtu.be/aaaaaaaaaaa"},
        ])
        self.db.remove_unavailable_dataset_item("ds:2")
        sheet = load_workbook(io.BytesIO(self.db.get_dataset_file_blob("ds"))).active
        self.assertEqual(sheet["C2"].value, "first")
        self.assertEqual(sheet.max_row, 2)


if __name__ == "__main__":
    unittest.main()
