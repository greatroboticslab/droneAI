"""Exercise the queue actions through the same HTTP routes as the dashboard."""
import io
import sys
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from openpyxl import Workbook, load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import app as gui  # noqa: E402
from db.db_store import DBStore  # noqa: E402
import validation_backend  # noqa: E402


class ValidationWorkflowTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.db = DBStore(str(Path(self.tmp.name) / "test.sqlite"))
        book = Workbook()
        book.active.append(["Persona Name", "Youtube Link"])
        book.active.append(["Ann", "https://youtu.be/aaaaaaaaaaa"])
        stream = io.BytesIO()
        book.save(stream)
        self.db.save_dataset("ds", "Flights", "flights.xlsx", stream.getvalue(), "tester")
        self.db.replace_dataset_items("ds", [{
            "item_key": "ds:1", "row_index": 1, "person_name": "Ann",
            "youtube_link": "https://youtu.be/aaaaaaaaaaa", "status": "in_progress",
            "locked_by": "labeler",
        }])
        self.db.upsert_validation_session(
            sid="session-1", youtube_link="https://youtu.be/aaaaaaaaaaa", status="running")
        db_patch = patch.object(gui, "db", self.db)
        db_patch.start()
        self.addCleanup(db_patch.stop)
        gui.app.config["TESTING"] = True
        self.client = gui.app.test_client()
        with self.client.session_transaction() as session:
            session["user"] = "labeler"
            session["current_item_key"] = "ds:1"
            session["current_dataset_key"] = "ds"
            session["current_validation_sid"] = "session-1"
            session["current_scenario_type"] = "Real flight"

    def test_done_cannot_label_queue_before_clips_are_final(self):
        with patch.object(gui, "get_validation_status", return_value={"state": "ready"}):
            response = self.client.post("/validation_release_lock")
        self.assertEqual(response.status_code, 409)
        self.assertEqual(self.db.get_dataset_item("ds:1")["status"], "in_progress")

        self.db.finalize_validation_session("session-1", 30.0, 1)
        with patch.object(gui, "get_validation_status", return_value={"state": "ready"}):
            response = self.client.post("/validation_release_lock")
        self.assertEqual(response.status_code, 200)
        self.assertEqual(self.db.get_dataset_item("ds:1")["status"], "labeled")
        sheet = load_workbook(io.BytesIO(self.db.get_dataset_file_blob("ds"))).active
        headers = [cell.value for cell in sheet[1]]
        self.assertEqual(sheet.cell(2, headers.index("Label Status") + 1).value, "Labeled")

    def test_remove_action_rejects_other_failures_then_removes_unavailable_video(self):
        with patch.object(gui, "get_validation_status", return_value={
            "state": "failed", "link": "https://youtu.be/aaaaaaaaaaa",
            "can_remove_unavailable": False,
        }):
            response = self.client.post("/validation_remove_unavailable")
        self.assertEqual(response.status_code, 409)
        self.assertIsNotNone(self.db.get_dataset_item("ds:1"))

        with patch.object(gui, "get_validation_status", return_value={
            "state": "failed", "link": "https://youtu.be/aaaaaaaaaaa",
            "can_remove_unavailable": True,
        }):
            response = self.client.post("/validation_remove_unavailable")
        self.assertEqual(response.status_code, 200)
        self.assertIsNone(self.db.get_dataset_item("ds:1"))
        sheet = load_workbook(io.BytesIO(self.db.get_dataset_file_blob("ds"))).active
        self.assertEqual(sheet.max_row, 1)
        download = self.client.get("/datasets/ds/download")
        self.assertEqual(download.status_code, 200)
        self.assertEqual(load_workbook(io.BytesIO(download.data)).active.max_row, 1)
        queue = self.client.get("/queue?dataset_key=ds")
        self.assertEqual(queue.status_code, 200)
        self.assertIn(b"Download labeling Excel", queue.data)

    def _ready_session(self, sid, at_sec):
        session = validation_backend.LabelingSession(sid)
        session.state = "ready"
        session.playback.set_current_time_sec(at_sec)
        validation_backend._register(session)
        self.addCleanup(validation_backend._sessions.pop, sid, None)
        return session

    def test_event_mark_requires_active_video_and_returns_saved_status(self):
        with patch.object(validation_backend, "db", self.db):
            self.assertIsNone(validation_backend.mark_event_now("no-such-session", "takeoff"))
            downloading = validation_backend.LabelingSession("session-1")
            validation_backend._register(downloading)
            self.addCleanup(validation_backend._sessions.pop, "session-1", None)
            self.assertIsNone(validation_backend.mark_event_now("session-1", "takeoff"))
            downloading.state = "ready"
            downloading.playback.set_current_time_sec(12.0)
            self.assertTrue(validation_backend.mark_event_now("session-1", "takeoff"))
        with self.db._conn() as conn:
            event = conn.execute("SELECT event_type, time_sec FROM validation_events").fetchone()
        self.assertEqual((event["event_type"], event["time_sec"]), ("takeoff", 12.0))

    def test_mark_returns_event_and_undo_removes_only_the_last_one(self):
        self._ready_session("session-1", 5.0)
        mark = validation_backend.mark_event_now
        with patch.object(validation_backend, "db", self.db):
            first = mark("session-1", "takeoff")
            second = mark("session-1", "land")
            self.assertEqual(first, {"index": 1, "type": "takeoff", "time_sec": 5.0})
            self.assertEqual(validation_backend.undo_last_event("session-1")["type"], "land")
            third = mark("session-1", "severe-crash")
        self.assertEqual(second["index"], 2)
        self.assertEqual(third["index"], 2)
        with self.db._conn() as conn:
            rows = conn.execute("SELECT idx, event_type FROM validation_events ORDER BY idx").fetchall()
        self.assertEqual([(r["idx"], r["event_type"]) for r in rows], [(1, "takeoff"), (2, "severe-crash")])

    def test_two_labelers_on_one_server_do_not_share_anything(self):
        ann = self._ready_session("sid-ann", 3.0)
        bob = self._ready_session("sid-bob", 40.0)
        with patch.object(validation_backend, "db", self.db):
            validation_backend.toggle_pause("sid-ann")
            validation_backend.skip_video("sid-bob", -5)
            validation_backend.mark_event_now("sid-ann", "takeoff", reviewer="ann")
            validation_backend.mark_event_now("sid-bob", "land", reviewer="bob")
            self.assertTrue(ann.playback.get_pause_flag())
            self.assertFalse(bob.playback.get_pause_flag())
            validation_backend.finish_validation_early("sid-ann")
            # Ann finishing does not stop Bob.
            self.assertTrue(validation_backend.mark_event_now("sid-bob", "minor-crash"))
        self.assertEqual(bob.playback.get_and_clear_skip_offset(), -5)
        self.assertEqual([e["type"] for e in validation_backend.list_current_events("sid-ann")], ["takeoff"])
        self.assertEqual([e["type"] for e in validation_backend.list_current_events("sid-bob")],
                         ["land", "minor-crash"])
        self.assertTrue(validation_backend.is_video_done("sid-ann"))
        self.assertFalse(validation_backend.is_video_done("sid-bob"))

    def test_routes_use_the_callers_own_session(self):
        self._ready_session("session-1", 7.0)
        self._ready_session("someone-else", 99.0)
        with patch.object(validation_backend, "db", self.db), patch.object(gui, "db", self.db):
            response = self.client.post("/mark_event", json={"event": "land"})
        self.assertEqual(response.get_json()["mark"]["time_sec"], 7.0)
        self.assertEqual(validation_backend.list_current_events("someone-else"), [])

    def test_unknown_session_status_is_idle(self):
        status = validation_backend.get_validation_status("gone")
        self.assertEqual(status["state"], "idle")

    def test_folder_names_are_never_shared(self):
        with tempfile.TemporaryDirectory() as root:
            a = validation_backend.get_unique_folder_name(root, "Ann")
            b = validation_backend.get_unique_folder_name(root, "Ann")
        self.assertNotEqual(a, b)

    def _add_items(self):
        self.db.replace_dataset_items("ds", [
            {"item_key": "ds:1", "row_index": 1, "person_name": "Ann",
             "youtube_link": "https://youtu.be/aaaaaaaaaaa", "status": "labeled"},
            {"item_key": "ds:2", "row_index": 2, "person_name": "Bob",
             "youtube_link": "https://youtu.be/bbbbbbbbbbb", "status": "not_labeled"},
            {"item_key": "ds:3", "row_index": 3, "person_name": "Cat",
             "youtube_link": "https://youtu.be/ccccccccccc", "status": "not_labeled",
             "source": "Real"},
        ])

    def test_queue_lists_labeled_videos_last(self):
        self._add_items()
        html = self.client.get("/queue?dataset_key=ds").get_data(as_text=True)
        self.assertLess(html.index("Bob"), html.index("Ann"))
        self.assertLess(html.index("Cat"), html.index("Ann"))
        self.assertIn("Label next video", html)

    def test_next_starts_first_unlabeled_and_skip_moves_on(self):
        self._add_items()
        with patch.object(gui, "start_validation_async", return_value="new-sid") as start, \
             patch.object(gui, "cancel_validation") as cancel, \
             patch.object(gui.mqtt_mgr, "publish_event"), patch.object(gui.mqtt_mgr, "publish_lock"):
            response = self.client.post("/queue/next", data={"dataset_key": "ds"})
            self.assertEqual(response.status_code, 302)
            self.assertIn("/validation_view_stream", response.headers["Location"])
            self.assertEqual(start.call_args.kwargs["youtube_link"], "https://youtu.be/bbbbbbbbbbb")
            self.assertEqual(self.db.get_dataset_item("ds:2")["status"], "in_progress")

            self.client.post("/queue/skip", data={"dataset_key": "ds"})
            cancel.assert_any_call("new-sid")   # the skipped video's session
            self.assertEqual(self.db.get_dataset_item("ds:2")["status"], "not_labeled")
            self.assertEqual(start.call_args.kwargs["youtube_link"], "https://youtu.be/ccccccccccc")
            self.assertEqual(start.call_args.kwargs["scenario_base"], "Real flight")
            with self.client.session_transaction() as session:
                self.assertEqual(session["current_item_key"], "ds:3")
                self.assertEqual(session["current_validation_sid"], "new-sid")

    def test_claim_is_exclusive(self):
        self._add_items()
        self.assertTrue(self.db.claim_dataset_item("ds:2", "ann"))
        self.assertFalse(self.db.claim_dataset_item("ds:2", "bob"))
        self.assertTrue(self.db.claim_dataset_item("ds:2", "ann"))   # resuming your own
        self.assertEqual(self.db.get_dataset_item("ds:2")["locked_by"], "ann")

    def test_next_skips_a_video_someone_claimed_a_moment_earlier(self):
        self._add_items()
        real_claim = self.db.claim_dataset_item

        def someone_else_was_faster(item_key, user, scenario=""):
            if item_key == "ds:2":
                real_claim(item_key, "bob")
            return real_claim(item_key, user, scenario)

        with patch.object(self.db, "claim_dataset_item", side_effect=someone_else_was_faster), \
             patch.object(gui, "start_validation_async", return_value="sid") as start, \
             patch.object(gui.mqtt_mgr, "publish_event"), patch.object(gui.mqtt_mgr, "publish_lock"):
            self.client.post("/queue/next", data={"dataset_key": "ds"})
        self.assertEqual(start.call_args.kwargs["youtube_link"], "https://youtu.be/ccccccccccc")
        self.assertEqual(self.db.get_dataset_item("ds:2")["locked_by"], "bob")
        self.assertEqual(self.db.get_dataset_item("ds:3")["locked_by"], "labeler")


class HoldAtEndTests(unittest.TestCase):
    def test_last_frame_is_held_until_stopped(self):
        import numpy as np
        import video_utils

        class FakeCapture:
            def __init__(self):
                self.left = 2
                self.pos = 0

            def read(self):
                if self.left == 0:
                    return False, None
                self.left -= 1
                self.pos += 1
                return True, np.zeros((4, 4, 3), dtype=np.uint8)

            def get(self, _prop):
                return self.pos

            def set(self, *_args):
                pass

            def release(self):
                raise AssertionError("held stream must not release the capture")

        video_utils.set_pause_flag(False)
        video_utils.get_and_clear_skip_offset()
        yielded = []
        stop_after = 6
        frames = video_utils.read_video_frames(
            FakeCapture(), 1000.0, hold_at_end=True,
            should_stop=lambda: len(yielded) >= stop_after)
        for frame in frames:
            yielded.append(frame)
        self.assertEqual(len(yielded), stop_after)
        self.assertEqual(yielded[-1], yielded[1])


if __name__ == "__main__":
    unittest.main()
