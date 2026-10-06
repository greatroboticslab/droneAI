"""Edit the uploaded dataset workbook when a queue row is finished or removed."""
from io import BytesIO

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter


PERSON_HEADERS = ("persona name", "person name", "person", "name")
LINK_HEADERS = ("youtube link", "youtube url", "youtube", "link", "url")
ITEM_KEY_HEADER = "DroneAI Queue ID"


def _column(sheet, names):
    for cell in sheet[1]:
        if str(cell.value or "").strip().lower() in names:
            return cell.column
    raise ValueError("The uploaded workbook no longer has its person and link columns.")


def _open_workbook(blob):
    try:
        return load_workbook(BytesIO(blob))
    except Exception:
        # The importer also accepts old .xls files. Convert these to .xlsx on
        # first edit; their original styling cannot be preserved by openpyxl.
        try:
            import pandas as pd
            frame = pd.read_excel(BytesIO(blob), sheet_name=0)
            stream = BytesIO()
            frame.to_excel(stream, index=False)
            return load_workbook(BytesIO(stream.getvalue()))
        except Exception as exc:
            raise ValueError("The uploaded Excel workbook could not be edited.") from exc


def edit_workbook(blob, item, all_items, *, remove=False, reset=False, labeled_by=""):
    """Return an .xlsx copy with one source row removed or marked labeled."""
    book = _open_workbook(blob)
    sheet = book.worksheets[0]
    person_col = _column(sheet, PERSON_HEADERS)
    link_col = _column(sheet, LINK_HEADERS)
    key_col = next((c.column for c in sheet[1] if c.value == ITEM_KEY_HEADER), None)
    if key_col is None:
        key_col = sheet.max_column + 1
        sheet.cell(1, key_col, ITEM_KEY_HEADER)
        by_value = {}
        for queued in sorted(all_items, key=lambda value: value["row_index"]):
            by_value.setdefault((queued["person_name"], queued["youtube_link"]), []).append(queued["item_key"])
        for row_number in range(2, sheet.max_row + 1):
            value = (str(sheet.cell(row_number, person_col).value or "").strip(),
                     str(sheet.cell(row_number, link_col).value or "").strip())
            keys = by_value.get(value)
            if keys:
                sheet.cell(row_number, key_col, keys.pop(0))
    sheet.column_dimensions[get_column_letter(key_col)].hidden = True
    status_col = next((c.column for c in sheet[1] if c.value == "Label Status"), None)
    candidates = [row for row in range(2, sheet.max_row + 1)
                  if sheet.cell(row, key_col).value == item["item_key"]]
    if not candidates:
        raise ValueError("The matching row was not found in the uploaded Excel workbook.")
    if reset and status_col is not None:
        row = next((r for r in candidates if sheet.cell(r, status_col).value == "Labeled"), candidates[0])
    else:
        row = next((r for r in candidates if status_col is None or
                    sheet.cell(r, status_col).value != "Labeled"), candidates[0])

    if remove:
        sheet.delete_rows(row)
    elif reset:
        if status_col is not None:
            sheet.cell(row, status_col).value = None
        by_col = next((c.column for c in sheet[1] if c.value == "Labeled By"), None)
        if by_col is not None:
            sheet.cell(row, by_col).value = None
    else:
        if status_col is None:
            status_col = sheet.max_column + 1
            sheet.cell(1, status_col, "Label Status")
        by_col = next((c.column for c in sheet[1] if c.value == "Labeled By"), None)
        if by_col is None:
            by_col = sheet.max_column + 1
            sheet.cell(1, by_col, "Labeled By")
        sheet.cell(row, status_col, "Labeled")
        sheet.cell(row, by_col, labeled_by)

    stream = BytesIO()
    book.save(stream)
    return stream.getvalue()
